# src/excel_writer.py
# Excel export functionality

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from datetime import datetime
import os
from .config import COLUMNS

def write_to_excel(rows, output_path=None):
    """
    Write extracted SLOs to an Excel file.
    
    Args:
        rows: List of dictionaries, each with keys matching COLUMNS
        output_path: Path to save the Excel file. If None, generates a timestamped filename.
    
    Returns:
        Path to the created Excel file.
    """
    if not rows:
        raise ValueError("No data to write to Excel")
    
    # Generate output path if not provided
    if output_path is None:
        timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
        output_path = f"data/output/slos_{timestamp}.xlsx"
    
    # Ensure directory exists
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "SLO Data"
    
    # Write headers
    for col_idx, header in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center')
        cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
    
    # Write data rows
    for row_idx, row_data in enumerate(rows, 2):
        for col_idx, column in enumerate(COLUMNS, 1):
            value = row_data.get(column, "")
            # Wrap long text
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    
    # Auto-adjust column widths
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)  # Cap at 50 characters
        ws.column_dimensions[column_letter].width = adjusted_width
    
    # Save the workbook
    wb.save(output_path)
    return output_path

def append_to_excel(rows, existing_file_path):
    """
    Append new rows to an existing Excel file.
    (Useful for batch processing multiple syllabi)
    """
    from openpyxl import load_workbook
    
    if not os.path.exists(existing_file_path):
        return write_to_excel(rows, existing_file_path)
    
    # Load existing workbook
    wb = load_workbook(existing_file_path)
    ws = wb.active
    
    # Find the next empty row
    next_row = ws.max_row + 1
    
    # Append data
    for row_data in rows:
        for col_idx, column in enumerate(COLUMNS, 1):
            ws.cell(row=next_row, column=col_idx, value=row_data.get(column, ""))
        next_row += 1
    
    wb.save(existing_file_path)
    return existing_file_path